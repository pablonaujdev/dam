# Original customization: Avannubo. License AGPL-3.

import base64
from datetime import date, datetime
from io import BytesIO
import math

import openpyxl
from psycopg2 import IntegrityError

from odoo import Command, fields, models, _
from odoo.exceptions import AccessError, UserError, ValidationError


class ImportInitialAccountMoveWiz(models.TransientModel):
    _name = 'import.initial.account.move.wiz'
    _description = 'Import Initial Account Move wiz'

    upload_file = fields.Binary(string='Upload File', help="Carga un Excel con las siete columnas de la plantilla de asientos.\nAdmite XLSX y XLSM; todos los asientos deben superar la validación.")
    file_name = fields.Char(string='File Name', help="Identifica el archivo Excel seleccionado para importar.\nEl nombre debe terminar en .xlsx o .xlsm.")
    date = fields.Date('Accounting Date', help="Determina la fecha contable de todos los asientos del archivo.\nLa columna FECHA del Excel se conserva como vencimiento de cada línea.")
    journal_id = fields.Many2one('account.journal', string='Journal', help="Selecciona el diario y la compañía de los asientos importados.\nDebe estar activo, ser accesible y trabajar en moneda de compañía.")
    account_dif_id = fields.Many2one('account.account', string='Offset Account', help="Conserva la selección de contrapartida del asistente original.\nCampo inactivo: los descuadres se rechazan sin añadir líneas automáticas.")

    def _initial_error_action(self, errors):
        wizard = self.env['error.log.wiz'].create({
            'message': _('There are errors in the import of the entries'),
            'error_line_ids': [Command.create({'error': error}) for error in errors],
        })
        return {
            'name': _('Errors during import'), 'type': 'ir.actions.act_window',
            'res_model': 'error.log.wiz', 'res_id': wizard.id,
            'view_mode': 'form', 'target': 'new',
        }

    def _initial_read_rows(self):
        if not self.upload_file:
            raise ValidationError(_('You must select a file.'))
        if not self.file_name or not self.file_name.lower().endswith(('.xlsx', '.xlsm')):
            raise ValidationError(_('Use an XLSX or XLSM file. Convert XLS files before importing.'))
        try:
            content = base64.b64decode(self.upload_file, validate=True)
            values_book = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
            try:
                formulas_book = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False)
                try:
                    sheet = values_book.worksheets[0]
                    formulas = formulas_book.worksheets[0]
                    headers = ('ASIENTO', 'FECHA', 'DEBIT', 'CREDIT', 'EMPRESA', 'CUENTA', 'DESCRIPCION')
                    actual = tuple(str(cell.value or '').strip().upper() for cell in next(sheet.iter_rows(max_row=1, max_col=7)))
                    if actual != headers:
                        raise ValidationError(_('The first seven columns must be: ASIENTO, FECHA, DEBIT, CREDIT, Empresa, CUENTA, DESCRIPCION.'))
                    rows = []
                    for number, (cached, original) in enumerate(zip(
                        sheet.iter_rows(min_row=2, max_col=7),
                        formulas.iter_rows(min_row=2, max_col=7),
                    ), start=2):
                        values = [cell.value for cell in cached]
                        uncached = [index for index, cell in enumerate(original)
                                    if cell.data_type == 'f' and cached[index].value is None]
                        if any(value is not None and value != '' for value in values) or uncached:
                            rows.append((number, values, uncached))
                    if not rows:
                        raise ValidationError(_('The file contains no entry lines.'))
                    return rows
                finally:
                    formulas_book.close()
            finally:
                values_book.close()
        except ValidationError:
            raise
        except Exception as error:
            raise ValidationError(_('Cannot read the Excel file. Save it as a valid XLSX or XLSM file.')) from error

    def _initial_identifier(self, value):
        if value is None or isinstance(value, bool):
            raise ValidationError(_('Entry identifier and account code must not be empty.'))
        if isinstance(value, float):
            if not math.isfinite(value):
                raise ValidationError(_('Identifiers must be text or finite numbers.'))
            value = int(value) if value.is_integer() else value
        text = str(value).strip()
        if not text:
            raise ValidationError(_('Entry identifier and account code must not be empty.'))
        if len(text) > 128:
            raise ValidationError(_('Identifiers and account codes must not exceed 128 characters.'))
        return text

    def _initial_due_date(self, value):
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if isinstance(value, str):
            for pattern in ('%Y-%m-%d', '%d/%m/%Y'):
                try:
                    return datetime.strptime(value.strip(), pattern).date()
                except ValueError:
                    continue
        raise ValidationError(_('FECHA must be an Excel date, YYYY-MM-DD or DD/MM/YYYY.'))

    def _initial_amount(self, value):
        if value is None or value == '':
            return 0.0
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            raise ValidationError(_('DEBIT and CREDIT must contain numeric amounts.'))
        amount = float(value)
        if not math.isfinite(amount) or amount < 0:
            raise ValidationError(_('DEBIT and CREDIT must be finite, non-negative amounts.'))
        return amount

    def _initial_prepare_entries(self, rows, company):
        accounts = self.env['account.account'].with_company(company).with_context(allowed_company_ids=[company.id])
        partners = self.env['res.partner'].with_company(company).with_context(allowed_company_ids=[company.id])
        account_cache, partner_cache, entries, errors = {}, {}, {}, []
        for number, values, uncached in rows:
            try:
                if uncached:
                    raise ValidationError(_('A formula has no calculated result. Recalculate and save the file in Excel.'))
                source = self._initial_identifier(values[0])
                code = self._initial_identifier(values[5])
                due_date = self._initial_due_date(values[1])
                debit, credit = self._initial_amount(values[2]), self._initial_amount(values[3])
                if debit > 0 and credit > 0:
                    raise ValidationError(_('A line cannot have both DEBIT and CREDIT greater than zero.'))
                if code not in account_cache:
                    account_cache[code] = accounts.search([
                        ('code', '=', code), ('company_ids', 'parent_of', company.id),
                    ])
                account = account_cache[code]
                if not account:
                    raise ValidationError(_('Active account %s was not found in the journal company.', code))
                if len(account) != 1:
                    raise ValidationError(_('Account code %s has multiple matches.', code))
                if account.currency_id and account.currency_id != company.currency_id:
                    raise ValidationError(_('Account %s requires a foreign currency unsupported by this template.', code))
                name = str(values[4]).strip() if values[4] is not None else ''
                partner = partners.browse()
                if name:
                    if name not in partner_cache:
                        partner_cache[name] = partners.search([
                            ('name', '=', name), ('company_id', 'in', [False, company.id]),
                        ])
                    partner = partner_cache[name]
                    if len(partner) != 1:
                        raise ValidationError(_('Contact %s must match exactly one active contact in this company.', name))
                elif account.account_type in ('asset_receivable', 'liability_payable'):
                    raise ValidationError(_('A contact name is required for receivable/payable account %s.', code))
                entry = entries.setdefault(source, {'rows': [], 'lines': [], 'debit': 0.0, 'credit': 0.0})
                entry['rows'].append(number)
                entry['debit'] += debit
                entry['credit'] += credit
                entry['lines'].append({
                    'account_id': account.id, 'partner_id': partner.id or False,
                    'date_maturity': due_date, 'debit': debit, 'credit': credit,
                    'name': str(values[6]) if values[6] not in (None, '') else _('Imported'),
                    # Entry imports never infer invoice taxes from account defaults.
                    'tax_ids': [Command.clear()],
                })
            except ValidationError as error:
                errors.append(_('Row %(row)s, entry %(entry)s: %(error)s',
                                row=number, entry=values[0] if values[0] is not None else '-', error=str(error)))
        for source, entry in entries.items():
            difference = entry['debit'] - entry['credit']
            if not math.isfinite(difference) or not company.currency_id.is_zero(difference):
                errors.append(_('Entry %(entry)s (rows %(rows)s) is not balanced: debit %(debit)s, credit %(credit)s.',
                                entry=source, rows=', '.join(map(str, entry['rows'])),
                                debit=entry['debit'], credit=entry['credit']))
        return entries, errors

    def import_xls_file(self):
        self.ensure_one()
        if not self.env.user.has_group('account.group_account_user') and not self.env.user.has_group('account.group_account_manager'):
            raise AccessError(_('Only accounting users can import entries.'))
        if not self.date or not self.journal_id:
            return self._initial_error_action([_('Select an accounting date and a journal.')])
        journal = self.journal_id
        journal.check_access('read')
        company = journal.company_id
        if company not in self.env.companies or not journal.active:
            return self._initial_error_action([_('Select an active journal in an allowed company.')])
        if journal.currency_id and journal.currency_id != company.currency_id:
            return self._initial_error_action([_('The journal must use the company currency.')])
        try:
            entries, errors = self._initial_prepare_entries(self._initial_read_rows(), company)
        except ValidationError as error:
            return self._initial_error_action([str(error)])
        moves_model = self.env['account.move'].with_company(company).with_context(allowed_company_ids=[company.id])
        for source, entry in entries.items():
            duplicate = moves_model.search([
                ('company_id', '=', company.id), ('journal_id', '=', journal.id), ('date', '=', self.date),
                '|', ('import_initial_asiento', '=', source),
                ('ref', '=', 'Import Accounting Entries %s' % source),
            ], limit=1)
            if duplicate:
                errors.append(_('Entry %(entry)s already exists for this company, journal and accounting date.', entry=source))
        if errors:
            return self._initial_error_action(errors)
        source, move_ids = '', []
        try:
            # One savepoint for the entire file: creation errors undo this import only.
            with self.env.cr.savepoint():
                for source, entry in entries.items():
                    move = moves_model.create({
                        'move_type': 'entry', 'state': 'draft', 'company_id': company.id,
                        'date': self.date, 'journal_id': journal.id,
                        'ref': 'Import Accounting Entries %s' % source,
                        'import_initial_asiento': source,
                        'line_ids': [Command.create(line) for line in entry['lines']],
                    })
                    move_ids.append(move.id)
        except IntegrityError as error:
            if 'import_initial_asiento_unique' in (error.diag.constraint_name or ''):
                message = _('The entry already exists, possibly from a concurrent import.')
            else:
                message = _('Odoo rejected an accounting value. Check the accounts, amounts and company configuration.')
            return self._initial_error_action([_('Entry %(entry)s: %(error)s', entry=source, error=message)])
        except (UserError, ValidationError) as error:
            return self._initial_error_action([_('Entry %(entry)s: %(error)s', entry=source, error=str(error))])
        return {
            'type': 'ir.actions.act_window', 'view_mode': 'list,form',
            'name': _('Account Move'), 'res_model': 'account.move',
            'domain': [('id', 'in', move_ids)], 'target': 'current',
        }
