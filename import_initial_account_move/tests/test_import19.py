import base64
from datetime import date
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import openpyxl
from psycopg2 import IntegrityError
from lxml import etree

from odoo import Command
from odoo.exceptions import AccessError, UserError
from odoo.tests import tagged
from odoo.addons.account.tests.common import AccountTestInvoicingCommon


@tagged('post_install', '-at_install', 'initial_import19')
class TestInitialImport19(AccountTestInvoicingCommon):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.account = cls.company_data['default_account_expense']
        cls.offset = cls.company_data['default_account_revenue']
        cls.journal = cls.company_data['default_journal_misc']
        cls.partner = cls.env['res.partner'].create({'name': 'MIAC Import Unique Contact'})
        cls.accounting_date = date(2026, 1, 1)
        cls.due = date(2026, 2, 15)

    def _rows(self, source='TEST-01', amount=100):
        return [
            [source, self.due, amount, 0, self.partner.name, self.account.code, 'Debit'],
            [source, self.due, 0, amount, None, self.offset.code, 'Credit'],
        ]

    def _wizard(self, rows=None, **extra):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(['ASIENTO', 'FECHA', 'DEBIT', 'CREDIT', 'Empresa', 'CUENTA', 'DESCRIPCION'])
        for row in rows if rows is not None else self._rows():
            sheet.append(row)
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()
        values = {'upload_file': base64.b64encode(buffer.getvalue()), 'file_name': 'opening.xlsx',
                  'date': self.accounting_date, 'journal_id': self.journal.id}
        values.update(extra)
        return self.env['import.initial.account.move.wiz'].create(values)

    def _moves(self, action):
        self.assertEqual(action['res_model'], 'account.move')
        return self.env['account.move'].search(action['domain'])

    def _errors(self, action):
        self.assertEqual(action['res_model'], 'error.log.wiz')
        wizard = self.env['error.log.wiz'].browse(action['res_id'])
        self.assertTrue(wizard.error_line_ids)
        return '\n'.join(wizard.error_line_ids.mapped('error'))

    def test_valid_grouping_last_row_blank_and_dates(self):
        first, second = self._rows('A'), self._rows('B', 25)
        rows = [first[0], second[0], [None] * 7, first[1], second[1]]
        moves = self._moves(self._wizard(rows).import_xls_file())
        self.assertEqual(len(moves), 2)
        self.assertEqual(set(moves.mapped('import_initial_asiento')), {'A', 'B'})
        self.assertTrue(all(move.state == 'draft' and move.move_type == 'entry' for move in moves))
        self.assertTrue(all(move.date == self.accounting_date for move in moves))
        self.assertTrue(all(line.date_maturity == self.due for line in moves.line_ids))
        self.assertEqual(sum(moves.line_ids.mapped('debit')), 125)
        self.assertEqual(sum(moves.line_ids.mapped('credit')), 125)
        self.assertEqual(moves.line_ids.filtered('debit').partner_id, self.partner)

    def test_dates_codes_and_identifiers(self):
        account = self.account.copy({'code': '00199001'})
        rows = self._rows(12.0)
        rows[0][1], rows[1][1] = '2026-02-15', '15/02/2026'
        rows[0][5] = '00199001'
        moves = self._moves(self._wizard(rows, file_name='opening.xlsm').import_xls_file())
        self.assertEqual(moves.import_initial_asiento, '12')
        self.assertEqual(moves.line_ids.filtered('debit').account_id, account)
        self.assertEqual(set(moves.line_ids.mapped('date_maturity')), {self.due})

    def test_validation_is_atomic_and_does_not_rollback_other_work(self):
        rows = self._rows('VALID') + self._rows('INVALID')
        rows[-1][5] = 'NO-SUCH-ACCOUNT'
        unrelated = self.env['res.partner'].create({'name': 'Must survive invalid import'})
        errors = self._errors(self._wizard(rows).import_xls_file())
        self.assertIn('Row 5', errors)
        self.assertTrue(unrelated.exists())
        self.assertFalse(self.env['account.move'].search([('import_initial_asiento', 'in', ['VALID', 'INVALID'])]))

    def test_creation_error_rolls_back_whole_file(self):
        wizard = self._wizard(self._rows('FIRST') + self._rows('SECOND'))
        unrelated = self.env['res.partner'].create({'name': 'Must survive creation error'})
        model = type(self.env['account.move'])
        original = model.create

        def reject_second(records, values):
            if isinstance(values, dict) and values.get('import_initial_asiento') == 'SECOND':
                raise UserError('Injected native creation error')
            return original(records, values)

        with patch.object(model, 'create', reject_second):
            errors = self._errors(wizard.import_xls_file())
        self.assertIn('SECOND', errors)
        self.assertTrue(unrelated.exists())
        self.assertFalse(self.env['account.move'].search([('import_initial_asiento', 'in', ['FIRST', 'SECOND'])]))

    def test_reject_unbalanced_and_inactive_offset(self):
        rows = self._rows()
        rows[1][3] = 50
        self.assertIn('not balanced', self._errors(self._wizard(rows, account_dif_id=self.offset.id).import_xls_file()))
        self.assertFalse(self.env['account.move'].search([('import_initial_asiento', '=', 'TEST-01')]))

    def test_invalid_amounts_dates_and_missing_ids(self):
        for column, value in ((0, None), (1, '31/02/2026'), (2, -1), (2, '100'), (2, True), (3, 1), (5, None)):
            with self.subTest(column=column, value=value):
                rows = self._rows()
                rows[0][column] = value
                self._errors(self._wizard(rows).import_xls_file())
        wizard = self._wizard()
        for value in (float('nan'), float('inf')):
            with self.assertRaises(UserError):
                wizard._initial_amount(value)

    def test_reject_formulas_without_cache_and_bad_files(self):
        rows = self._rows()
        rows[0][2] = '=50+50'
        self.assertIn('formula', self._errors(self._wizard(rows).import_xls_file()))
        self._errors(self._wizard(upload_file=base64.b64encode(b'broken zip')).import_xls_file())
        self._errors(self._wizard(file_name='opening.xls').import_xls_file())
        self._errors(self._wizard(upload_file=False).import_xls_file())
        self._errors(self._wizard([]).import_xls_file())

    def test_reject_wrong_headers(self):
        wizard = self._wizard()
        workbook = openpyxl.load_workbook(BytesIO(base64.b64decode(wizard.upload_file)))
        workbook.active['A1'] = 'UNKNOWN'
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()
        wizard.upload_file = base64.b64encode(buffer.getvalue())
        self.assertIn('first seven columns', self._errors(wizard.import_xls_file()))

    def test_partner_exact_unique_required_and_active(self):
        duplicate = self.partner.copy({'name': self.partner.name})
        self.assertIn('exactly one', self._errors(self._wizard().import_xls_file()))
        duplicate.active = False
        self._moves(self._wizard().import_xls_file())
        rows = self._rows('MISSING')
        rows[0][4] = 'Non-existent name'
        self._errors(self._wizard(rows).import_xls_file())
        rows[0][4] = None
        rows[0][5] = self.company_data['default_account_receivable'].code
        self.assertIn('required', self._errors(self._wizard(rows).import_xls_file()))
        rows[0][4] = self.partner.name
        self._moves(self._wizard(rows).import_xls_file())

    def test_accounts_companies_and_currency(self):
        inactive = self.account.copy({'code': '990881', 'active': False})
        rows = self._rows()
        rows[0][5] = inactive.with_context(active_test=False).code
        self._errors(self._wizard(rows).import_xls_file())
        other_company = self.env['res.company'].create({'name': 'Other Import Company'})
        self.partner.company_id = other_company
        self._errors(self._wizard().import_xls_file())
        self.partner.company_id = False
        foreign = self.env.ref('base.EUR') if self.env.company.currency_id != self.env.ref('base.EUR') else self.env.ref('base.USD')
        self.journal.currency_id = foreign
        self.assertIn('currency', self._errors(self._wizard().import_xls_file()))
        self.journal.currency_id = False
        self.account.currency_id = foreign
        self.assertIn('foreign currency', self._errors(self._wizard().import_xls_file()))

    def test_new_and_legacy_duplicates_and_other_dates_journals(self):
        wizard = self._wizard()
        move = self._moves(wizard.import_xls_file())
        self.assertIn('already exists', self._errors(wizard.import_xls_file()))
        legacy = self._moves(self._wizard(self._rows('LEGACY')).import_xls_file())
        legacy.import_initial_asiento = False
        self.assertIn('already exists', self._errors(self._wizard(self._rows('LEGACY')).import_xls_file()))
        self._moves(self._wizard(date=date(2026, 1, 2)).import_xls_file())
        journal = self.journal.copy({'name': 'Other Import Journal', 'code': 'IMT'})
        self._moves(self._wizard(journal_id=journal.id).import_xls_file())
        copied = move.copy()
        self.assertFalse(copied.import_initial_asiento)
        with self.assertRaises(IntegrityError), self.cr.savepoint():
            move.copy({'import_initial_asiento': move.import_initial_asiento,
                       'date': move.date, 'journal_id': move.journal_id.id})

    def test_permissions_and_shared_error_wizard(self):
        from odoo.tests.common import new_test_user
        user = new_test_user(self.env, login='initial_import_basic', groups='base.group_user')
        with self.assertRaises(AccessError):
            self.env['import.initial.account.move.wiz'].with_user(user).create({})
        shared = self.env['error.log.wiz'].with_user(user).create({
            'message': 'Shared error', 'error_line_ids': [Command.create({'error': 'Other importer'})]})
        self.assertEqual(shared.error_line_ids.error, 'Other importer')
        self.assertEqual(shared.error_line_ids.wiz_id, shared)

    def test_same_identifier_in_other_company(self):
        self._moves(self._wizard().import_xls_file())
        company = self.env['res.company'].create({'name': 'Independent Import Company'})
        account_model = self.env['account.account'].with_company(company)
        account_model.create([
            {'name': 'Import debit', 'code': self.account.code, 'account_type': 'expense',
             'company_ids': [Command.set(company.ids)]},
            {'name': 'Import credit', 'code': self.offset.code, 'account_type': 'income',
             'company_ids': [Command.set(company.ids)]},
        ])
        journal = self.env['account.journal'].create({'name': 'Independent Import', 'code': 'ICP',
            'type': 'general', 'company_id': company.id})
        wizard = self._wizard(journal_id=journal.id).with_company(company)
        action = wizard.import_xls_file()
        self.assertEqual(action['res_model'], 'account.move')
        moves = self.env['account.move'].with_company(company).search(action['domain'])
        self.assertEqual(moves.company_id, company)

    def test_template_and_views(self):
        path = Path(__file__).parents[1] / 'static/data/Import_Account_Move.xlsx'
        workbook = openpyxl.load_workbook(path, data_only=True)
        rows = list(workbook.active.values)
        workbook.close()
        self.assertEqual(rows[0], ('ASIENTO', 'FECHA', 'DEBIT', 'CREDIT', 'Empresa', 'CUENTA', 'DESCRIPCION'))
        self.assertGreaterEqual(len(rows), 3)
        self.assertEqual(sum(row[2] or 0 for row in rows[1:]), sum(row[3] or 0 for row in rows[1:]))
        self.assertTrue(all(not (row[2] and row[3]) for row in rows[1:]))
        view = self.env.ref('error_line_wiz.error_log_wiz_view_form')
        arch = etree.fromstring(view.arch_db.encode())
        self.assertTrue(arch.xpath('//list'))
        self.assertFalse(arch.xpath('//tree'))
        self.assertTrue(self.env.ref('import_initial_account_move.action_importar_account_move'))
